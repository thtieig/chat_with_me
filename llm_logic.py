# --- my_chat_app/llm_logic.py ---
import os
import sys
from typing import List, Dict, Generator, Union
from dotenv import load_dotenv
from openai import OpenAI, APIError
import pandas as pd
import docx
import PyPDF2
import json
import csv
from openpyxl import load_workbook
from zipfile import ZipFile, BadZipFile
import chardet
import logging
import bleach
import yaml
import io # Keep io for type hints, even if not used in placeholder process_files

load_dotenv(override=True)

# Configure basic logging
logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(levelname)s - %(message)s')

def load_config(config_file="config.yaml"):
    """Loads configuration from a YAML file."""
    try:
        with open(config_file, 'r') as f:
            config_data = yaml.safe_load(f)
            if not config_data or 'providers' not in config_data or 'personas' not in config_data:
                raise ValueError("Invalid config structure. Missing 'providers' or 'personas'.")
            logging.info(f"Configuration loaded successfully from {config_file}")
            return config_data
    except FileNotFoundError:
        logging.error(f"Configuration file not found: {config_file}")
        sys.exit(1) # Exit if config is essential and missing
    except (yaml.YAMLError, ValueError) as e:
        logging.error(f"Error parsing configuration file '{config_file}': {e}")
        sys.exit(1) # Exit on bad config

config = load_config()

# --- Client Configuration ---

def configure_ionos_client(provider_config):
    """Configures and returns an OpenAI client for IONOS."""
    api_key_env = provider_config.get('api_key_env')
    api_key = os.getenv(api_key_env) if api_key_env else None
    base_url = provider_config.get('base_url')
    if not api_key:
        raise ValueError(f"API key environment variable '{api_key_env}' not found for IONOS.")
    if not base_url:
         raise ValueError("Base URL not configured for IONOS.")
    logging.info(f"Configuring IONOS client with base URL: {base_url}")
    return OpenAI(api_key=api_key, base_url=base_url)

def configure_google_client(provider_config):
    """Configures and returns an OpenAI client for Google (OpenAI-compatible endpoint)."""
    api_key_env = provider_config.get('api_key_env')
    api_key = os.getenv(api_key_env) if api_key_env else None
    base_url = provider_config.get('base_url')
    if not api_key:
        raise ValueError(f"API key environment variable '{api_key_env}' not found for Google.")
    if not base_url:
         raise ValueError("Base URL not configured for Google.")
    logging.info(f"Configuring Google client (OpenAI endpoint) with base URL: {base_url}")
    return OpenAI(api_key=api_key, base_url=base_url)

def configure_ollama_client(provider_config):
    """Configures and returns an OpenAI client for Ollama."""
    # Use OLLAMA_BASE_URL from .env first, then config.yaml, then default
    default_ollama_url = 'http://localhost:11434/v1'
    base_url = os.getenv("OLLAMA_BASE_URL", provider_config.get('base_url', default_ollama_url))
    api_key = provider_config.get('api_key', 'ollama') # Default key for Ollama often 'ollama'

    if not base_url:
        # This case should ideally not be hit due to defaults, but added for safety
        logging.warning(f"Ollama base URL not explicitly configured, using default: {default_ollama_url}")
        base_url = default_ollama_url

    logging.info(f"Configuring Ollama client with base URL: {base_url}")
    # Ollama doesn't typically require a real API key when accessed locally
    return OpenAI(base_url=base_url, api_key=api_key)

# --- Persona and Sanitization ---

def get_persona_message(persona_name: str, config_data: dict) -> str:
    """Retrieves the system prompt for a given persona name from config."""
    if not config_data or 'personas' not in config_data:
        logging.warning("Personas configuration is missing or invalid. Using default.")
        return "You are a helpful assistant."

    persona_messages = config_data.get('personas', {})
    # Find the first persona listed as the default if 'Default' doesn't exist
    default_persona_name = next(iter(persona_messages.keys()), None) if persona_messages else None
    default_message = "You are a helpful assistant."
    if default_persona_name:
        default_message = persona_messages.get(default_persona_name, default_message)

    # Get the message for the requested persona, fall back to 'Default', then the first one, then the hardcoded default
    message = persona_messages.get(persona_name, persona_messages.get('Default', default_message))
    logging.info(f"Using persona '{persona_name}': {message[:80]}...") # Log snippet
    return message

def detect_encoding(file_data: bytes) -> str:
    """Detects the encoding of file data using chardet."""
    result = chardet.detect(file_data)
    encoding = result['encoding'] if result['encoding'] else 'utf-8'
    logging.debug(f"Detected encoding: {encoding} with confidence {result['confidence']}")
    return encoding


def sanitize_html(html_content: str) -> str:
    """Sanitizes HTML content based on settings in config.yaml."""
    if not config or 'html_sanitization' not in config:
        logging.warning("HTML sanitization config not found. Returning content unsanitized.")
        return html_content # Return original if no config

    settings = config.get('html_sanitization', {})
    allowed_tags = settings.get('allowed_tags', [])
    allowed_attributes = settings.get('allowed_attributes', {})
    logging.debug(f"Sanitizing HTML with tags: {allowed_tags}, attributes: {allowed_attributes.keys()}")
    # strip=True removes disallowed tags entirely
    # strip_comments=True is default and good practice
    return bleach.clean(html_content, tags=allowed_tags, attributes=allowed_attributes, strip=True)

# --- File Processing (Placeholders - Needs Implementation) ---

def process_text_file(file_stream: io.BytesIO, filename: str) -> str:
    """Placeholder: Processes a text-based file."""
    logging.info(f"Processing text file (placeholder): {filename}")
    try:
        raw_data = file_stream.read()
        encoding = detect_encoding(raw_data)
        content = raw_data.decode(encoding, errors='replace')
        return f"\n--- Content from {filename} ---\n{content}\n---\n"
    except Exception as e:
        logging.error(f"Error processing text file {filename}: {e}")
        return f"\n--- Error processing {filename}: Could not read content ---\n"

def process_pdf_file(file_stream: io.BytesIO, filename: str) -> str:
    """Placeholder: Processes a PDF file."""
    logging.info(f"Processing PDF file (placeholder): {filename}")
    # Add PyPDF2 logic here
    content = f"\n--- Content from PDF {filename} (Processing not implemented) ---\n"
    return content

def process_docx_file(file_stream: io.BytesIO, filename: str) -> str:
    """Placeholder: Processes a DOCX file."""
    logging.info(f"Processing DOCX file (placeholder): {filename}")
    # Add python-docx logic here
    content = f"\n--- Content from DOCX {filename} (Processing not implemented) ---\n"
    return content

# Add placeholders for other file types (CSV, Excel, JSON, ZIP etc.) as needed

def process_files(files: List[tuple[io.BytesIO, str]]) -> str:
    """
    Processes a list of uploaded files based on their type.
    Args:
        files: A list of tuples, where each tuple contains (file_stream, filename).
               Using werkzeug FileStorage directly might be better if Flask context is available.
               Here assuming BytesIO stream and filename are passed.
    Returns:
        A single string concatenating the processed content of all files.
    """
    all_content = ""
    if not files:
        return ""

    logging.info(f"Processing {len(files)} files...")
    for file_stream, filename in files:
        try:
            # Reset stream position in case it was read before
            file_stream.seek(0)

            # Basic type checking based on filename extension
            _, extension = os.path.splitext(filename.lower())

            if extension in ['.txt', '.py', '.js', '.css', '.html', '.md', '.log', '.yaml', '.xml', '.json', '.csv']: 
                content = process_text_file(file_stream, filename)
            elif extension == '.pdf':
                content = process_pdf_file(file_stream, filename)
            elif extension == '.docx':
                content = process_docx_file(file_stream, filename)
            # Add elif conditions for other supported types (Excel, Zip, etc.)
            else:
                logging.warning(f"Skipping unsupported file type: {filename}")

            all_content += content

        except Exception as e:
            logging.error(f"Failed to process file '{filename}': {e}", exc_info=True)
            all_content += f"\n--- Unexpected error processing file {filename} ---\n"
        finally:
            pass  # Be cautious with closing request streams directly

    logging.info(f"Finished processing files. Total content length (approx): {len(all_content)}")
    return all_content


# --- Core LLM Streaming Logic ---

def stream_question(
    provider: str,
    model: str,
    persona: str,
    history: List[Dict[str, str]],
    prompt: str,
    # files: List[io.BytesIO], # Change type if using Flask's FileStorage directly
    files: List[tuple[io.BytesIO, str]], # Use tuple (stream, filename)
    config_data: dict
) -> Generator[str, None, None]:
    """
    Gets response from the AI model as a stream of text chunks.
    Handles client selection, message preparation, file processing, and API call.
    Yields content chunks or error messages.
    """
    client = None
    try:
        # --- 1. Get Provider Configuration ---
        provider_config = config_data.get('providers', {}).get(provider)
        if not provider_config:
            raise ValueError(f"Configuration for provider '{provider}' not found.")

        # --- 2. Configure AI Client ---
        if provider == 'IONOS':
            client = configure_ionos_client(provider_config)
        elif provider == 'GOOGLE':
            # Assuming Google uses an OpenAI-compatible endpoint as configured
            client = configure_google_client(provider_config)
        elif provider == 'OLLAMA':
            client = configure_ollama_client(provider_config)
        else:
            raise ValueError(f"Unsupported provider selected: {provider}")

        if not client:
            # This should ideally be caught by specific config errors, but as a safeguard:
            raise ConnectionError(f"Failed to initialize client for provider {provider}.")

        # --- 3. Prepare Messages for API ---
        system_message_content = get_persona_message(persona, config_data)
        messages = [{"role": "system", "content": system_message_content}]

        # Add history (filter out placeholders if frontend sends them)
        valid_history = [
            msg for msg in history
            if isinstance(msg, dict) and msg.get("role") and msg.get("content") and msg.get("content") != '...'
        ]
        messages.extend(valid_history)

        # --- 4. Process Files and Combine with Prompt ---
        # This calls the (currently placeholder) file processing logic
        file_context = process_files(files) if files else ""

        # Combine user prompt with file context
        final_prompt = prompt
        if file_context:
            # Add a separator and the extracted file content
            final_prompt = f"{prompt}\n\n--- Context from Uploaded Files ---\n{file_context}\n--- End of File Context ---"
            logging.info(f"Added context from {len(files)} files to the prompt.")

        messages.append({"role": "user", "content": final_prompt})

        logging.info(f"Sending request to {provider} model {model}. Total messages: {len(messages)}")
        # Optional: Log message content snippets for debugging (be mindful of size/privacy)
        # logging.debug(f"System Prompt: {messages[0]['content'][:100]}...")
        # logging.debug(f"User Prompt: {messages[-1]['content'][:200]}...")

        # --- 5. Make Streaming API Call ---
        stream = client.chat.completions.create(
            model=model,
            messages=messages,
            stream=True
        )

        # --- 6. Process Stream and Yield Chunks ---
        for response_chunk in stream:
            # The exact structure of 'response_chunk' can vary slightly,
            # but for OpenAI compatible APIs, it's usually like this:
            if response_chunk.choices:
                delta = response_chunk.choices[0].delta
                content_chunk = delta.content
                if content_chunk:  # Check if there is actual text content
                    #logging.debug(f"Yielding chunk: '{content_chunk}'") # Very verbose
                    yield content_chunk # <-- This is where the defined 'content_chunk' is yielded

    # --- Error Handling ---
    except APIError as e:
        error_message = f"API Error from {provider}: Status={e.status_code}, Message={getattr(e, 'message', str(e))}"
        logging.error(error_message, exc_info=True)
        yield f"Error: {error_message}" # Yield error to frontend
    except ValueError as e: # Catch config/value errors (e.g., missing keys, unsupported provider)
        error_message = f"Configuration or Value Error: {str(e)}"
        logging.error(error_message, exc_info=False) # No need for full traceback for config issues
        yield f"Error: {error_message}"
    except ConnectionError as e: # Catch client connection issues
         error_message = f"Connection Error for {provider}: {str(e)}"
         logging.error(error_message, exc_info=True)
         yield f"Error: {error_message}"
    except Exception as e: # Catch any other unexpected errors during streaming
        error_message = f"An unexpected error occurred during generation: {str(e)}"
        logging.error(error_message, exc_info=True)
        yield f"Error: {error_message}"
    finally:
        # Close the client? Typically not needed for standard OpenAI client usage pattern.
        # If using custom clients with manual session management, cleanup might go here.
        logging.info(f"Stream generation finished or terminated for provider {provider}.")